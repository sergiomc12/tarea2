"""
build_docx.py -- genera report/main.docx replicando (en la medida en que
Word lo permite) el mismo contenido y formato IEEE de report/main.tex.

IMPORTANTE: el notebook, report/main.tex y este script son tres artefactos
separados -- nada se sincroniza solo entre ellos. Este script es la UNICA
forma de actualizar el .docx, porque un binario no se puede editar a mano
con las herramientas de este proyecto. Cada vez que se porte un paso nuevo
del notebook a report/main.tex, hay que portarlo TAMBIEN aqui y volver a
correr `python build_docx.py` para regenerar main.docx.

Uso:
    uv run python report/build_docx.py
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent
FIG_TARGET = HERE / ".." / "output" / "figures" / "fig01_target_balance.png"
OUT_PATH = HERE / "main.docx"

BODY_FONT = "Times New Roman"
CODE_FONT = "Courier New"


# ---------------------------------------------------------------------
# Helpers de bajo nivel (python-docx no expone columnas ni bordes de
# parrafo en su API de alto nivel -- se manipula el XML directamente)
# ---------------------------------------------------------------------

def set_columns(section, n):
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sectPr.append(cols)
    cols.set(qn('w:num'), str(n))
    cols.set(qn('w:space'), '432')  # ~0.3in de calle entre columnas


def add_continuous_section(doc, columns=2):
    new_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(new_section, columns)
    return new_section


def set_paragraph_box_border(paragraph, size=8, color="000000"):
    """Encuadra un parrafo completo (usado para 'cada funcion en su
    propio recuadro', requisito literal de hw02-2026-2.docx)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:space'), '4')
        el.set(qn('w:color'), color)
        pBdr.append(el)
    pPr.append(pBdr)


def set_base_font(doc, name=BODY_FONT, size=10):
    style = doc.styles['Normal']
    style.font.name = name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def add_code_function(doc, title, lines):
    """Una funcion = un recuadro independiente, Courier New 8pt --
    requisito de la pregunta de codigo fuente de hw02-2026-2.docx."""
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(line if line else " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(8)
        rpr = run._element.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), CODE_FONT)
        rFonts.set(qn('w:hAnsi'), CODE_FONT)
        rpr.append(rFonts)
        set_paragraph_box_border(p)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(title)
    run.italic = True
    run.font.size = Pt(9)
    doc.add_paragraph()  # espaciador


def add_body_paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.2)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


def add_heading_roman(doc, number, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{number}. {text.upper()}")
    run.bold = True
    run.font.size = Pt(10)
    return p


def add_subheading(doc, letter, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"{letter}. {text}")
    run.bold = True
    run.italic = True
    run.font.size = Pt(10)
    return p


def add_pending(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(f"[Pendiente de desarrollo -- {text}]")
    run.italic = True
    return p


def set_table_fixed_layout(table, total_width_in):
    """python-docx no fija el layout por defecto -- sin esto, Word/
    LibreOffice autoajustan las columnas y pueden colapsar/solapar
    contenido (bug real observado y corregido durante el desarrollo)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(int(total_width_in * 1440)))
    tblW.set(qn('w:type'), 'dxa')


def set_row_cell_widths(row, widths_in):
    for cell, w in zip(row.cells, widths_in):
        cell.width = Inches(w)


def add_table_ieee(doc, roman, caption, headers, rows, col_widths_in=None):
    """Tabla con numeracion romana y titulo ARRIBA en versalitas (Small
    Caps), como exige el formato IEEE. `col_widths_in` es
    obligatorio en la practica -- sin ancho fijo por columna, Word/
    LibreOffice colapsan columnas con contenido corto (ver
    set_table_fixed_layout)."""
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(f"TABLE {roman}. {caption.upper()}")
    run.bold = False
    run.font.small_caps = True
    run.font.size = Pt(9)

    n = len(headers)
    if col_widths_in is None:
        col_widths_in = [3.2 / n] * n
    total_width_in = sum(col_widths_in)

    table = doc.add_table(rows=1, cols=n)
    table.style = 'Table Grid'
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_fixed_layout(table, total_width_in)

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].paragraphs[0].add_run(h).bold = True
        for r in hdr_cells[i].paragraphs[0].runs:
            r.font.size = Pt(9)
    set_row_cell_widths(table.rows[0], col_widths_in)

    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9)
        set_row_cell_widths(table.rows[-1], col_widths_in)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def add_figure_ieee(doc, path, number, caption, width_in=3.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(f"Fig. {number}. {caption}")
    run.font.size = Pt(9)
    doc.add_paragraph()


# ---------------------------------------------------------------------
# Documento
# ---------------------------------------------------------------------

def mainProcess():
    doc = Document()
    set_base_font(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = section.right_margin = Inches(0.75)
    section.top_margin = section.bottom_margin = Inches(0.75)

    # --- Titulo, autor y filiacion (una columna, como en la plantilla IEEE) ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "Análisis Exploratorio, Limpieza, Codificación, Balanceo y "
        "Ranking de Atributos sobre el Dataset Home Credit Default Risk"
    )
    run.bold = True
    run.font.size = Pt(20)

    for line in ["Mondragon Contreras, Sergio", "Curso de Machine Learning",
                 "Tarea 2 -- 31 de julio de 2026"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line).font.size = Pt(11)

    doc.add_paragraph()

    # --- Cuerpo a dos columnas ---
    add_continuous_section(doc, columns=2)

    p = doc.add_paragraph()
    r = p.add_run("Abstract"); r.bold = True; r.italic = True
    p.add_run("—Resumen pendiente de completar.").italic = True
    p2 = doc.add_paragraph()
    r2 = p2.add_run("Index Terms"); r2.bold = True; r2.italic = True
    p2.add_run("—Home Credit Default Risk, EDA, calidad de datos, "
                "balanceo de clases, selección de atributos, PCA.").italic = True

    add_heading_roman(doc, "I", "Introducción")
    add_body_paragraph(
        doc,
        "Este informe desarrolla la Tarea 2 del curso de Machine Learning: "
        "un análisis exploratorio, de calidad, limpieza, codificación, "
        "balanceo y selección de atributos sobre el dataset público de la "
        "competencia Kaggle Home Credit Default Risk. El desarrollo "
        "completo (código, resultados intermedios y figuras) se produjo "
        "en un notebook Jupyter (notebooks/tarea2_home_credit.ipynb) y se "
        "resume en este documento."
    )
    add_body_paragraph(
        doc,
        "las secciones IV-A, IV-B y IV-C (pasos 1 a 3 del procedimiento) "
        "están desarrolladas con resultados reales, obtenidos ejecutando "
        "el notebook contra el archivo de datos completo. Las secciones "
        "IV-D a IV-H (pasos 4 a 8) están en desarrollo activo.",
        bold_lead="Estado de este informe: ",
    )

    add_heading_roman(doc, "II", "El Problema")
    add_body_paragraph(
        doc,
        "Se cuenta con datos reales de solicitudes de crédito de la "
        "entidad Home Credit, incluyendo si cada solicitante terminó "
        "incumpliendo con su pago (columna TARGET). El problema, en "
        "términos propios, consiste en explorar a fondo ese conjunto de "
        "datos, medir su calidad, corregir sus defectos, transformarlo a "
        "una forma que un algoritmo de aprendizaje pueda consumir, y "
        "determinar qué atributos concentran la mayor capacidad para "
        "anticipar si un cliente va a incumplir. Le importa a Home Credit "
        "porque su modelo de negocio depende de extender crédito "
        "responsablemente a personas sin historial crediticio tradicional."
    )

    add_heading_roman(doc, "III", "Metodología de Trabajo")
    add_body_paragraph(
        doc,
        "El procedimiento sigue ocho pasos: (1) descripción del problema, "
        "(2) descripción general del dataset, (3) análisis del TARGET, "
        "(4) calidad de los datos, (5) limpieza de datos, (6) conversión "
        "de categóricas a numéricas, (7) balanceo de clases, y (8) "
        "ranking de atributos por capacidad predictiva (técnicas Filter) "
        "con histogramas, dispersión y PCA. Los pasos 4 y 5 están "
        "explícitamente ponderados con mayor peso en la evaluación."
    )

    add_heading_roman(doc, "IV", "Desarrollo de la Solución")

    add_subheading(doc, "A", "Descripción del problema")
    add_body_paragraph(
        doc,
        "El archivo application_train.csv proviene de la competencia "
        "pública de Kaggle Home Credit Default Risk (2018). Home Credit "
        "es una institución financiera no bancaria fundada en 1997 en la "
        "República Checa, que opera en 14 países y se especializa en "
        "otorgar préstamos a personas con historial crediticio limitado "
        "o inexistente. Al año 2018, el grupo reportaba más de 29 "
        "millones de clientes y más de 160 millones de préstamos "
        "otorgados, la mayoría en Asia."
    )
    add_body_paragraph(
        doc,
        "Cada fila de application_train.csv es una solicitud de préstamo "
        "(SK_ID_CURR), con datos estáticos del solicitante y la columna "
        "TARGET (1 = pago tardío significativo en los primeros plazos, "
        "0 = cualquier otro caso). La métrica de evaluación original de "
        "la competencia era ROC-AUC -- relevante porque el dataset está "
        "desbalanceado (Sección IV-C), lo que hace del accuracy una "
        "métrica engañosa."
    )
    add_body_paragraph(
        doc,
        "el dataset completo de Kaggle incluye 10 archivos; esta tarea, "
        "por instrucción explícita del enunciado, usa solo "
        "application_train.csv.",
        bold_lead="Nota de alcance: ",
    )

    add_subheading(doc, "B", "Descripción general del dataset")
    add_body_paragraph(
        doc,
        "El dataset tiene 307,511 filas y 122 columnas. La Tabla I "
        "resume la composición por tipo de dato."
    )
    add_table_ieee(
        doc, "I", "Tipos de dato en application_train.csv",
        ["Tipo", "N. de columnas"],
        [["Numérico (int64/float64)", "106"],
         ["Categórico (object/str)", "16"]],
        col_widths_in=[2.3, 1.0],
    )
    add_body_paragraph(
        doc,
        "Dentro de las 106 columnas numéricas hay una mezcla heterogénea: "
        "identificadores, indicadores binarios 0/1, conteos enteros, "
        "montos continuos y variables de tiempo relativas en días "
        "negativos (DAYS_BIRTH, DAYS_EMPLOYED). Entre las 16 categóricas, "
        "la cardinalidad es muy dispar: la mayoría tiene entre 2 y 8 "
        "categorías, mientras que ORGANIZATION_TYPE tiene 58."
    )
    add_body_paragraph(
        doc,
        "307,511 filas es un volumen amplio para entrenar un modelo "
        "robusto y da margen a las técnicas de balanceo (Sección IV-G). "
        "Sin embargo, el volumen total no resuelve el problema de "
        "desbalance de clases (Sección IV-C).",
        bold_lead="Influencia en las secciones posteriores: ",
    )

    add_subheading(doc, "C", "Análisis del Target")
    add_body_paragraph(
        doc,
        "La Tabla II y la Fig. 1 resumen la distribución de la variable "
        "objetivo."
    )
    add_table_ieee(
        doc, "II", "Distribución de la variable TARGET",
        ["Clase", "Descripción", "% del total"],
        [["0", "No incumple (repaga a tiempo)", "91.93%"],
         ["1", "Incumple (pago tardío significativo)", "8.07%"]],
        col_widths_in=[0.5, 1.9, 0.8],
    )
    if FIG_TARGET.exists():
        add_figure_ieee(doc, FIG_TARGET, 1,
                         "Conteo y proporción de solicitudes por clase TARGET.")
    add_body_paragraph(
        doc,
        "La clase TARGET=0 representa 282,686 solicitudes (91.93%) y la "
        "clase TARGET=1 representa 24,825 solicitudes (8.07%) -- una "
        "razón de aproximadamente 11.4:1. El criterio operativo de este "
        "informe es que una clase minoritaria por debajo del 10-20% ya "
        "constituye desbalance; aquí está en 8.07%, por debajo incluso "
        "del límite más estricto."
    )
    add_body_paragraph(
        doc,
        "el conjunto de técnicas que corrigen la disparidad de ejemplos "
        "entre clases para que el modelo no optimice casi exclusivamente "
        "para la clase mayoritaria. Sin balanceo, un modelo puede "
        "alcanzar ~92% de accuracy prediciendo siempre TARGET=0 -- inútil "
        "en la práctica. El desbalance típicamente hunde el recall y el "
        "F1-score de la clase minoritaria, razón por la que la "
        "competencia original evalúa con ROC-AUC.",
        bold_lead="Qué es el balanceo de datos: ",
    )
    add_body_paragraph(
        doc,
        "en riesgo crediticio el costo de los dos tipos de error no es "
        "simétrico: un falso negativo (dejar pasar a quien incumplirá) "
        "cuesta dinero real; un falso positivo (rechazar a quien sí "
        "habría pagado) es solo costo de oportunidad. Esto sugiere "
        "preferir una técnica de balanceo que privilegie el recall de "
        "TARGET=1 sobre la precisión general (retomado en Sección IV-G).",
        bold_lead="Postura propia: ",
    )

    for letter, title_txt, detail in [
        ("D", "Calidad de los Datos",
         "paso 4. % de nulos por columna, filas con nulos, outliers, "
         "duplicados, columnas de una sola categoría, primer EDA (Anexo 1)."),
        ("E", "Limpieza de Datos",
         "paso 5. Tratamiento del valor centinela 365243 en DAYS_EMPLOYED "
         "y criterio de retención de EXT_SOURCE_1 pese a 56% de nulos."),
        ("F", "Conversión de Atributos Categóricos en Numéricos",
         "paso 6. Técnicas del Anexo 2 según el tipo de cada columna."),
        ("G", "Balanceo de Datos",
         "paso 7. Técnica candidata: SMOTE o class_weight='balanced'."),
        ("H", "Ranking de Atributos según su Capacidad Predictiva",
         "paso 8, condicionado al paso 7. Técnicas Filter, histogramas "
         "por clase, dispersión, PCA original vs. seleccionado."),
    ]:
        add_subheading(doc, letter, title_txt)
        add_pending(doc, detail)

    add_heading_roman(doc, "V", "Resultados")
    add_pending(doc, "se completa cuando cierren los 8 pasos del desarrollo.")

    add_heading_roman(doc, "VI", "Conclusiones")
    add_pending(
        doc,
        "con indicadores numéricos, no opiniones. Hasta ahora, "
        "verificado: el dataset tiene 307,511x122 (Sección IV-B) y un "
        "desbalance de clases de 11.4:1 que requiere corrección explícita "
        "antes de modelar (Sección IV-C)."
    )

    add_heading_roman(doc, "VII", "Recomendaciones")
    add_pending(doc, "mínimo 3, con análisis de viabilidad, al cierre del desarrollo.")

    # --- Codigo fuente: una columna (requisito literal del enunciado) ---
    add_continuous_section(doc, columns=1)
    add_heading_roman(doc, "VIII", "Código Fuente")
    p = doc.add_paragraph()
    p.add_run(
        "Código tomado y refactorizado del desarrollo real en "
        "notebooks/tarea2_home_credit.ipynb, ya ejecutado y verificado "
        "contra el dataset completo (Secciones IV-B y IV-C)."
    ).italic = True

    add_code_function(doc, "Listing 1: Clase HomeCreditAnalysis -- un método por paso del procedimiento", [
        'class HomeCreditAnalysis:',
        '    """Orquesta el analisis completo de application_train.csv',
        '    (Tarea 2). Cada metodo publico implementa un paso del',
        '    procedimiento exigido por hw02-2026-2.docx. El estado se',
        '    guarda como atributos de instancia para que los pasos',
        '    posteriores lo reutilicen sin recibirlo/devolverlo por',
        '    parametro."""',
        '',
        '    def __init__(self, data_path="../data/application_train.csv",',
        '                 figures_dir="../output/figures",',
        '                 tables_dir="../output/tables", dpi=300):',
        '        self.data_path = data_path',
        '        self.figures_dir = figures_dir',
        '        self.tables_dir = tables_dir',
        '        self.dpi = dpi',
        '        self.df = None',
        '        self.resumen_tipos = None',
        '        self.cardinalidad = None',
        '        self.numeric_summary = None',
        '        self.resumen_target = None',
        '',
        '    def cargar_datos(self):',
        '        """Paso 1-2 (carga): unico archivo en el alcance."""',
        '        self.df = pd.read_csv(self.data_path)',
        '        print(f"application_train.csv -> filas: "',
        '              f"{self.df.shape[0]:,}  columnas: {self.df.shape[1]}")',
        '        return self.df',
        '',
        '    def describir_dataset(self):',
        '        """Paso 2: tipos de dato, cardinalidad, Tabla I."""',
        '        df = self.df',
        '        # pandas 3.x reporta columnas de texto como dtype "str"',
        '        # (antes "object") -- se cuentan ambas etiquetas.',
        '        dtype_counts = df.dtypes.astype(str).value_counts()',
        '        n_numericas = int(dtype_counts.get("int64", 0)',
        '                           + dtype_counts.get("float64", 0))',
        '        n_categoricas = int(dtype_counts.get("object", 0)',
        '                             + dtype_counts.get("str", 0))',
        '',
        '        cat_cols = df.select_dtypes(include="object").columns.tolist()',
        '        self.cardinalidad = (df[cat_cols].nunique()',
        '                              .sort_values(ascending=False))',
        '',
        '        self.numeric_summary = (df.select_dtypes(include=[np.number])',
        '                                 .describe().T)',
        '        self.numeric_summary["rango"] = (self.numeric_summary["max"]',
        '                                          - self.numeric_summary["min"])',
        '',
        '        self.resumen_tipos = pd.DataFrame({',
        '            "tipo": ["numerico (int64/float64)",',
        '                     "categorico (object/str)"],',
        '            "n_columnas": [n_numericas, n_categoricas],',
        '        })',
        '        self.resumen_tipos.to_csv(',
        '            f"{self.tables_dir}/tabla01_tipos_de_dato.csv", index=False)',
        '        return self.resumen_tipos, self.cardinalidad',
        '',
        '    def analizar_target(self):',
        '        """Paso 3: balance de clases, Tabla II y Fig. 1."""',
        '        df = self.df',
        '        target_counts = df["TARGET"].value_counts().sort_index()',
        '        target_pct = (target_counts / target_counts.sum()',
        '                      * 100).round(2)',
        '        self.resumen_target = pd.DataFrame({',
        '            "clase": target_counts.index,',
        '            "conteo": target_counts.values,',
        '            "porcentaje": target_pct.values,',
        '        })',
        '        self.resumen_target.to_csv(',
        '            f"{self.tables_dir}/tabla02_balance_target.csv",',
        '            index=False)',
        '',
        '        fig, axes = plt.subplots(1, 2, figsize=(11, 5))',
        '        sns.barplot(x=self.resumen_target["clase"].astype(str),',
        '                    y=self.resumen_target["conteo"], ax=axes[0],',
        '                    color="#4C72B0")',
        '        axes[0].set_title(',
        '            "Fig. 1a. Conteo de solicitudes por clase TARGET")',
        '',
        '        etiquetas = [f"TARGET={c}\\n{p}%" for c, p in',
        '                     zip(self.resumen_target["clase"],',
        '                         self.resumen_target["porcentaje"])]',
        '        axes[1].pie(self.resumen_target["conteo"], labels=etiquetas,',
        '                    colors=["#4C72B0", "#DD8452"], startangle=90)',
        '        axes[1].set_title("Fig. 1b. Proporcion de clases TARGET")',
        '',
        '        plt.tight_layout()',
        '        plt.savefig(f"{self.figures_dir}/fig01_target_balance.png",',
        '                    dpi=self.dpi, bbox_inches="tight")',
        '        plt.close(fig)',
        '        return self.resumen_target',
        '',
        '    # TODO paso 4: calidad_de_los_datos(self)',
        '    # TODO paso 5: limpiar_datos(self)',
        '    # TODO paso 6: convertir_categoricas(self)',
        '    # TODO paso 7: balancear_clases(self)',
        '    # TODO paso 8: ranking_atributos(self)',
    ])

    add_code_function(doc, "Listing 2: mainProcess -- orquestador principal del análisis", [
        'def mainProcess():',
        '    """',
        '    Punto de entrada unico del analisis (Tarea 2).',
        '    Pasos 1-3 implementados y verificados contra el dataset',
        '    completo (307,511 x 122). Pasos 4-8 (calidad, limpieza,',
        '    encoding, balanceo, ranking+PCA) en desarrollo -- ver',
        '    hw02-2026-2.docx, seccion PROCEDIMIENTO.',
        '    """',
        '    analysis = HomeCreditAnalysis()',
        '    analysis.cargar_datos()',
        '    analysis.describir_dataset()',
        '    analysis.analizar_target()',
        '',
        '    # TODO paso 4: analysis.calidad_de_los_datos()',
        '    # TODO paso 5: analysis.limpiar_datos()',
        '    # TODO paso 6: analysis.convertir_categoricas()',
        '    # TODO paso 7: analysis.balancear_clases()',
        '    # TODO paso 8: analysis.ranking_atributos()',
        '',
        '    return analysis',
        '',
        '',
        'if __name__ == "__main__":',
        '    mainProcess()',
    ])

    # --- Referencias: vuelve a dos columnas ---
    add_continuous_section(doc, columns=2)
    ref_heading = doc.add_paragraph()
    ref_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ref_heading.paragraph_format.space_before = Pt(12)
    ref_heading.paragraph_format.space_after = Pt(6)
    ref_run = ref_heading.add_run("REFERENCIAS")
    ref_run.bold = True
    ref_run.font.size = Pt(10)
    refs = [
        "[1] M. Kuhn and K. Johnson, Feature Engineering and Selection: A "
        "Practical Approach for Predictive Models. Boca Raton, FL, USA: "
        "Chapman & Hall/CRC Press, 2019.",
        "[2] H. He and Y. Ma, Imbalanced Learning: Foundations, Algorithms "
        "and Applications. Hoboken, NJ, USA: Wiley-IEEE Press, 2013.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.add_run(ref).font.size = Pt(9)

    doc.save(str(OUT_PATH))
    print(f"Generado: {OUT_PATH}")


if __name__ == "__main__":
    mainProcess()
